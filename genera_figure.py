#!/usr/bin/env python3
"""
genera_figure.py - genera le Figure 1-7 della tesi dal workbook di benchmark.

Uso:
    python3 genera_figure.py Benchmark_grafici.xlsx
    python3 genera_figure.py Benchmark_grafici.xlsx --out fig --formato png --dpi 200
    python3 genera_figure.py Benchmark_grafici.xlsx --formato pdf

Legge i fogli "Data Set 1".."Data Set 4" (le colonne sono individuate per nome
di intestazione) e produce, nella cartella di uscita:

    fig1_errore_relativo            errore relativo medio per profilo
    fig2_errore_assoluto            errore assoluto medio (36 configurazioni)
    fig3_traiettorie_3srv           traiettorie con bande, configurazione a 3 server
    fig4_diagnostica                S/N e f1/f2 (mediana) per profilo
    fig5_intorno                    intorno [Chao log inf., Jackknife sup.] a t = 30
    fig6_accumulazione              curve di accumulazione di S
    fig7_traiettorie_<profilo>      traiettorie di tutte le configurazioni
    tab_spettro_N<N>.tex            tabella LaTeX dello spettro delle frequenze
                                    per una singola configurazione (default:
                                    1000 host e 3 server, cioe' N = 1003)

Opzioni per la tabella dello spettro:
    --spettro-host 1000 --spettro-server 3 --spettro-giorni 1 3 5 10

Convenzioni (le stesse descritte nella tesi):
  - le stime di Huggins superiori a SOGLIA_ESCLUSIONE * N sono escluse dalle
    medie delle Figure 1 e 2;
  - f1/f2 e' riportato come mediana sulle configurazioni del profilo;
  - nelle traiettorie l'asse verticale e' limitato a LIMITE_ASSE * N e i valori
    fuori scala sono segnalati da un triangolo con il valore.

Dipendenze: pip install openpyxl matplotlib
"""
import argparse
import os
import statistics as st
import sys
from collections import defaultdict

try:
    import openpyxl
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
except ImportError:
    sys.exit("Servono openpyxl e matplotlib:  pip install openpyxl matplotlib")

FOGLI = ["Data Set 1", "Data Set 2", "Data Set 3", "Data Set 4"]
PROFILI = ("molto_frequente", "bilanciato", "poco_frequente")
ETICHETTA = {"molto_frequente": "molto frequente", "bilanciato": "bilanciato",
             "poco_frequente": "poco frequente"}
TAGLIE = (1000, 500, 100)
SERVER = (3, 10, 15, 30)
GIORNI = list(range(1, 11))

SOGLIA_ESCLUSIONE = 3.0   # Huggins > 3N escluso dalle medie (Fig. 1-2)
LIMITE_ASSE = 1.6         # asse verticale delle traiettorie limitato a 1,6 N

# nome di intestazione nel workbook -> nome interno
COLONNE = {
    "profilo_benchmark": "prof", "giorno_cattura": "g", "S_obs": "S",
    "f1": "f1", "f2": "f2", "f3": "f3", "f4": "f4", "f5": "f5", "f_più_alte": "fhi",
    "stima_jackknife": "J", "stima_chao": "Ch", "stima_huggings": "H",
    "conf_jackknife_lower": "Jlo", "conf_jackknife_higher": "Jhi",
    "conf_chao_lower": "Chlo", "conf_chao_upper": "Chhi",
    "conf_chao_log_lower": "Cllo", "conf_chao_log_upper": "Clhi",
    "conf_huggings_lower": "Hlo", "conf_huggings_upper": "Hhi",
    "popolazione_reale": "N",
}

STILE = {"J": dict(color="tab:blue", marker="o", ls="-", label="Jackknife"),
         "Ch": dict(color="tab:red", marker="s", ls="--", label="Chao"),
         "H": dict(color="tab:green", marker="^", ls="-.", label="Huggins")}
COLORE_PROFILO = {"molto_frequente": "tab:blue", "bilanciato": "tab:orange",
                  "poco_frequente": "tab:red"}


# --------------------------------------------------------------------------- #
# lettura
# --------------------------------------------------------------------------- #
def leggi(xlsx):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    righe = []
    for nome in FOGLI:
        if nome not in wb.sheetnames:
            sys.exit("ERRORE: foglio '%s' assente nel workbook." % nome)
        ws = wb[nome]
        idx = None
        for r in ws.iter_rows(values_only=True):
            if not r:
                continue
            if r[0] == "profilo_benchmark":
                idx = {COLONNE[h]: i for i, h in enumerate(r) if h in COLONNE}
                continue
            if idx is None or r[0] not in PROFILI:
                continue
            d = {k: r[i] for k, i in idx.items()}
            if not isinstance(d.get("N"), (int, float)) or not isinstance(d.get("J"), (int, float)):
                continue
            N = int(d["N"])
            d["host"] = max(h for h in TAGLIE if h <= N)
            d["srv"] = N - d["host"]
            righe.append(d)
    if not righe:
        sys.exit("ERRORE: nessuna riga di dati trovata.")
    return righe


def media(v):
    v = [x for x in v if x is not None]
    return sum(v) / len(v) if v else float("nan")


def err(d, k):
    return 100.0 * (d[k] - d["N"]) / d["N"]


def err_filtrato(d, k):
    if k == "H" and d["H"] > SOGLIA_ESCLUSIONE * d["N"]:
        return None
    return err(d, k)


# --------------------------------------------------------------------------- #
# figure
# --------------------------------------------------------------------------- #
def fig1(R, out):
    fig, ax = plt.subplots(1, 3, figsize=(13, 4.6), sharey=True)
    for a, p in zip(ax, PROFILI):
        for k in STILE:
            y = [media(err_filtrato(d, k) for d in R if d["g"] == g and d["prof"] == p) for g in GIORNI]
            a.plot(GIORNI, y, ms=5, **STILE[k])
        a.axhline(0, color="k", lw=0.8)
        a.set_title("profilo " + ETICHETTA[p]); a.set_xticks(GIORNI)
        a.set_xlabel("giornata di rilevazione"); a.grid(alpha=.3)
    ax[0].set_ylabel(r"errore relativo medio $(\hat N-N)/N$  [%]")
    ax[0].legend(loc="lower right")
    fig.suptitle("Errore relativo medio al crescere dello sforzo di campionamento "
                 "(media su 12 configurazioni per profilo)")
    fig.tight_layout(); out(fig, "fig1_errore_relativo")


def fig2(R, out):
    mae = {k: [media(abs(e) if e is not None else None
                     for e in (err_filtrato(d, k) for d in R if d["g"] == g)) for g in GIORNI]
           for k in STILE}
    # prima giornata in cui Jackknife diventa piu' accurato di Chao
    inv = next((g for g, j, c in zip(GIORNI, mae["J"], mae["Ch"]) if j < c), None)
    fig, a = plt.subplots(figsize=(6.4, 4.8))
    for k in STILE:
        a.plot(GIORNI, mae[k], lw=2, ms=6, **STILE[k])
    if inv:
        a.axvline(inv, color="grey", ls=":")
        a.annotate("inversione\ndell'ordinamento", xy=(inv, mae["J"][inv - 1]),
                   xytext=(inv + 1.2, max(mae["J"]) * 0.7), color="grey", fontsize=9,
                   arrowprops=dict(arrowstyle="->", color="grey"))
    a.set_xticks(GIORNI); a.set_ylim(0, None); a.grid(alpha=.3)
    a.set_xlabel("giornata di rilevazione"); a.set_ylabel("errore relativo assoluto medio  [%]")
    a.set_title("Complementarità fra Jackknife e Chao\n(media sulle 36 configurazioni)")
    a.legend(); fig.tight_layout(); out(fig, "fig2_errore_assoluto")


def traiettoria(a, serie, bande):
    N = serie[1]["N"]
    a.axhline(N, color="k", lw=1.3, label="$N$ reale")
    a.plot(GIORNI, [serie[g]["S"] for g in GIORNI], color="grey", ls=":", label="$S$ osservati")
    tetto = LIMITE_ASSE * N
    pts = [serie[g]["S"] for g in GIORNI] + [N] + \
          [serie[g][k] for g in GIORNI for k in STILE if serie[g][k] <= tetto]
    lo, hi = min(pts), max(pts)
    pad = 0.08 * (hi - lo)
    if bande:
        for k, (l, h) in (("J", ("Jlo", "Jhi")), ("Ch", ("Cllo", "Clhi")), ("H", ("Hlo", "Hhi"))):
            a.fill_between(GIORNI, [serie[g][l] for g in GIORNI], [serie[g][h] for g in GIORNI],
                           color=STILE[k]["color"], alpha=.12, lw=0)
    for k in STILE:
        y = [serie[g][k] if serie[g][k] <= tetto else float("nan") for g in GIORNI]
        a.plot(GIORNI, y, ms=4, **STILE[k])
        for g in GIORNI:
            if serie[g][k] > tetto:
                a.plot([g], [hi + pad * 0.2], marker="^", color=STILE[k]["color"], ms=6)
                a.annotate("%.0f" % serie[g][k], xy=(g, hi + pad * 0.6), ha="center",
                           fontsize=7, color=STILE[k]["color"])
    a.set_ylim(lo - pad, hi + pad * 1.4); a.set_xticks(GIORNI); a.grid(alpha=.3)


def fig3(T, out):
    fig, ax = plt.subplots(3, 3, figsize=(13, 11))
    for i, h in enumerate(TAGLIE):
        for j, p in enumerate(PROFILI):
            a = ax[i][j]
            traiettoria(a, T[(3, p, h)], True)
            if i == 0: a.set_title("profilo " + ETICHETTA[p])
            if j == 0: a.set_ylabel("$N=%d$\nnumerosità stimata" % (h + 3))
            if i == 2: a.set_xlabel("giornata di rilevazione")
    ax[0][0].legend(fontsize=8, loc="lower right")
    fig.suptitle("Traiettorie delle stime con intervallo di fiducia al 95% "
                 "(3 server; banda di Chao log-trasformata)")
    fig.tight_layout(rect=(0, 0, 1, .97)); out(fig, "fig3_traiettorie_3srv")


def fig4(R, out):
    fig, ax = plt.subplots(1, 2, figsize=(12, 4.6))
    for p in PROFILI:
        c = COLORE_PROFILO[p]
        ax[0].plot(GIORNI, [100 * media(d["S"] / d["N"] for d in R if d["g"] == g and d["prof"] == p)
                            for g in GIORNI], marker="o", color=c, label=ETICHETTA[p])
        ax[1].plot(GIORNI, [st.median(d["f1"] / d["f2"] for d in R
                                      if d["g"] == g and d["prof"] == p and d["f2"])
                            for g in GIORNI], marker="o", color=c, label=ETICHETTA[p])
    ax[0].set_ylabel("copertura campionaria $S/N$  [%]")
    ax[0].set_title("Frazione di popolazione osservata")
    ax[1].axhline(1, color="k", ls="--", lw=1)
    ax[1].set_ylabel("$f_1/f_2$ (mediana)")
    ax[1].set_title("Grado di sottocampionamento\n(indicatore osservabile senza conoscere $N$)")
    for a in ax:
        a.set_xticks(GIORNI); a.grid(alpha=.3); a.legend(); a.set_xlabel("giornata di rilevazione")
    fig.tight_layout(); out(fig, "fig4_diagnostica")


def fig5(R, out):
    rs = sorted((d for d in R if d["g"] == 10),
                key=lambda d: (d["host"], d["srv"], PROFILI.index(d["prof"])))
    fig, a = plt.subplots(figsize=(10, 10))
    dentro = 0
    for y, d in enumerate(rs):
        N = d["N"]
        lo, hi = 100 * (d["Cllo"] - N) / N, 100 * (d["Jhi"] - N) / N
        ok = d["Cllo"] <= N <= d["Jhi"]; dentro += ok
        a.plot([lo, hi], [y, y], color="tab:green" if ok else "tab:red", lw=3)
        a.plot(err(d, "Ch"), y, "s", color="tab:red")
        a.plot(err(d, "J"), y, "o", color="tab:blue")
    a.axvline(0, color="k"); a.set_yticks(range(len(rs)))
    a.set_yticklabels(["$N=%d$, %s, %d srv" % (d["N"], ETICHETTA[d["prof"]], d["srv"]) for d in rs],
                      fontsize=8)
    a.set_xlabel("scarto da $N$  [%]"); a.grid(axis="x", alpha=.3)
    a.set_title("Intorno [Chao log inf., Jackknife sup.] a $t=30$\n"
                "verde: contiene $N$ (%d su %d).  Quadrato: $\\hat N_{Ch}$, cerchio: $\\hat N_J$"
                % (dentro, len(rs)))
    fig.tight_layout(); out(fig, "fig5_intorno")


def fig6(T, out):
    fig, ax = plt.subplots(3, 4, figsize=(15, 9.5))
    for i, h in enumerate(TAGLIE):
        for j, s in enumerate(SERVER):
            a = ax[i][j]
            for p in PROFILI:
                serie = T[(s, p, h)]
                a.plot(GIORNI, [serie[g]["S"] for g in GIORNI], marker="o", ms=4,
                       color=COLORE_PROFILO[p], label=ETICHETTA[p])
            a.axhline(h + s, color="k", label="$N$ reale")
            a.set_ylim(0, (h + s) * 1.08); a.set_xticks(GIORNI); a.grid(alpha=.3)
            if i == 0: a.set_title("%d server" % s)
            if j == 0: a.set_ylabel("$N\\approx%d$\nnodi osservati $S$" % h)
            if i == 2: a.set_xlabel("giornata")
    ax[0][0].legend(fontsize=8, loc="lower right")
    fig.suptitle("Curve di accumulazione delle catture: nodi distinti osservati "
                 "al crescere dello sforzo di campionamento")
    fig.tight_layout(rect=(0, 0, 1, .97)); out(fig, "fig6_accumulazione")


def fig7(T, out):
    for p in PROFILI:
        fig, ax = plt.subplots(3, 4, figsize=(15, 9.5))
        for i, h in enumerate(TAGLIE):
            for j, s in enumerate(SERVER):
                a = ax[i][j]
                traiettoria(a, T[(s, p, h)], False)
                if i == 0: a.set_title("$N=%d$ (%d server)" % (h + s, s), fontsize=10)
                if j == 0: a.set_ylabel("numerosità stimata")
                if i == 2: a.set_xlabel("giornata")
        ax[0][0].legend(fontsize=7, loc="lower right")
        fig.suptitle("Traiettorie delle stime, profilo %s — tutte le configurazioni" % ETICHETTA[p])
        fig.tight_layout(rect=(0, 0, 1, .97)); out(fig, "fig7_traiettorie_%s" % p)


def tabella_spettro(T, host, srv, giorni, cartella):
    """Tabella LaTeX: quota f_k/S per classe e copertura S/N, per una sola
    configurazione (host, srv), nelle giornate indicate, per i tre profili."""
    N = host + srv
    classi = ("f1", "f2", "f3", "f4", "f5", "fhi")
    num = lambda v: ("%.1f" % v).replace(".", "{,}")
    righe = []
    for i, g in enumerate(giorni):
        if i:
            righe.append("\\midrule")
        for p in PROFILI:
            d = T[(srv, p, host)][g]
            quote = " & ".join("$%s$" % num(100.0 * d[c] / d["S"]) for c in classi)
            righe.append("%2d & %-15s & %s & $%s\\%%$ \\\\"
                         % (g, ETICHETTA[p], quote, num(100.0 * d["S"] / d["N"])))
    testo = r"""\begin{center}
\captionof{table}{Ripartizione percentuale dei nodi osservati fra le classi di
    frequenza nella configurazione con %d host e %d server ($N = %d$), la stessa
    della Figura~\ref{fig:spettro}. Ogni valore è la quota $f_k/S$ rispetto al
    numero $S$ di nodi osservati; l'ultima colonna è la copertura $S/N$. Alla
    prima giornata vale $t = 3$, cosicché le classi $k \geq 4$ sono
    necessariamente vuote.}
\label{tab:spettro}
\small
\begin{tabular}{@{}cl rrrrrr r@{}}
\toprule
giorno & profilo & $f_1$ & $f_2$ & $f_3$ & $f_4$ & $f_5$ & $f_{>5}$ & $S/N$ \\
\midrule
%s
\bottomrule
\end{tabular}
\end{center}
""" % (host, srv, N, "\n".join(righe))
    percorso = os.path.join(cartella, "tab_spettro_N%d.tex" % N)
    with open(percorso, "w", encoding="utf-8") as fh:
        fh.write(testo)
    return percorso


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="workbook di benchmark (es. Benchmark_grafici.xlsx)")
    ap.add_argument("--out", default="fig", help="cartella di uscita (default: fig)")
    ap.add_argument("--formato", default="png", choices=["png", "pdf", "svg", "jpg"])
    ap.add_argument("--dpi", type=int, default=200)
    ap.add_argument("--spettro-host", type=int, default=1000, choices=TAGLIE)
    ap.add_argument("--spettro-server", type=int, default=3, choices=SERVER)
    ap.add_argument("--spettro-giorni", type=int, nargs="+", default=[1, 3, 5, 10])
    a = ap.parse_args()

    R = leggi(a.xlsx)
    T = defaultdict(dict)
    for d in R:
        T[(d["srv"], d["prof"], d["host"])][d["g"]] = d
    mancanti = [(s, p, h) for s in SERVER for p in PROFILI for h in TAGLIE
                if len(T[(s, p, h)]) < len(GIORNI)]
    if mancanti:
        sys.exit("ERRORE: configurazioni incomplete nel workbook: %s" % mancanti)

    os.makedirs(a.out, exist_ok=True)
    scritti = []

    def out(fig, nome):
        percorso = os.path.join(a.out, "%s.%s" % (nome, a.formato))
        fig.savefig(percorso, dpi=a.dpi)
        plt.close(fig)
        scritti.append(percorso)

    fig1(R, out); fig2(R, out); fig3(T, out); fig4(R, out)
    fig5(R, out); fig6(T, out); fig7(T, out)
    scritti.append(tabella_spettro(T, a.spettro_host, a.spettro_server,
                                   a.spettro_giorni, a.out))

    esclusi = [(d["N"], d["prof"], d["g"], round(d["H"])) for d in R
               if d["H"] > SOGLIA_ESCLUSIONE * d["N"]]
    print("File generati (%d):" % len(scritti))
    for s in scritti:
        print("  " + s)
    if esclusi:
        print("Stime di Huggins > %gN escluse dalle medie (N, profilo, giornata, stima):"
              % SOGLIA_ESCLUSIONE)
        for e in esclusi:
            print("  ", e)


if __name__ == "__main__":
    main()