#!/usr/bin/env python3
"""
rigenera_grafici.py - ricostruisce i grafici del workbook di benchmark.

Da lanciare DOPO bench_to_excel.py, sullo stesso file:

    python3 bench_to_excel.py --logs logs --xlsx Benchmark_grafici.xlsx
    python3 rigenera_grafici.py Benchmark_grafici.xlsx

Parte dai tre grafici del Data Set 1 presenti su ciascun foglio grafici e li
replica per ogni altro Data Set che abbia le stime compilate. E' idempotente:
prima di ricostruire scarta i grafici generati dall'esecuzione precedente,
quindi puo' essere rilanciato quante volte si vuole.

Cosa fa, oltre a replicare:
  - titoli uniformi "Descrizione (N nodi + S server) - Data Set X", con il
    numero di server dedotto da popolazione_reale (colonna X)
  - stesso ordine di serie in tutti i grafici, quindi stessi colori
  - riferimenti di serie, nomi e categorie ricalcolati dal blocco giusto
  - titoli degli assi e griglia su entrambi gli assi
  - valori memorizzati nei grafici riallineati alle celle

Opzioni:
    --out FILE     scrive altrove invece di sovrascrivere
    --no-pulizia   non tocca la colonna Z ne' la legenda colori
"""

import argparse
import os
import re
import shutil
import sys
import tempfile
import zipfile

try:
    import openpyxl
    from openpyxl.utils import column_index_from_string as col_index
except ImportError:
    sys.exit("Serve openpyxl:  pip install openpyxl")

# --- layout del foglio dati --------------------------------------------------
GRUPPI = [2, 34, 66]          # prima riga di ogni gruppo (1000, 500, 100 nodi)
RIGHE_BLOCCO = 10
PROFILI = ['molto_frequente', 'bilanciato', 'poco_frequente']
COL_POP = 'X'
COL_GIORNO = 'B'

# ordine canonico delle serie: nome -> colonna. L'ordine e' anche quello dei
# colori, perche' Excel assegna accent1..accent5 in base alla posizione.
SERIE_TRAFFICO = [('S_obs', 'D'), ('stima_jackknife', 'K'),
                  ('stima_chao', 'L'), ('stima_huggings', 'M')]
SERIE_ACCUMULO = [('catture_%s' % p, 'D') for p in PROFILI]

ASSI = {'b': 'Giorni di cattura', 'l': 'Numero di nodi'}
# modello di titolo asse, preso da un grafico che ce l'ha gia' formattato:
# alcuni grafici del template hanno il titolo asse vuoto o assente.
MODELLO_ASSE = {}
PASSO_RIGHE = 22              # altezza di una riga di grafici, in righe Excel


# --- utilita' XML ------------------------------------------------------------
def blocchi(xml, tag):
    return list(re.finditer(r'<c:%s>.*?</c:%s>' % (tag, tag), xml, flags=re.S))


def esc(t):
    return t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def testo_singolo(blocco, testo):
    """Riduce il rich text di un <c:tx> a un solo run che contiene `testo`."""
    i, j = blocco.find('<c:tx>'), blocco.find('</c:tx>')
    if i < 0:
        return blocco
    j += len('</c:tx>')
    tx = blocco[i:j]
    run = list(re.finditer(r'<a:r>.*?</a:r>', tx, flags=re.S))
    if not run:
        return blocco
    primo = run[0]
    nuovo = re.sub(r'<a:t>.*?</a:t>', '<a:t>%s</a:t>' % esc(testo),
                   primo.group(0), flags=re.S)
    tx2 = tx[:primo.start()] + nuovo + tx[primo.end():]
    for r in run[1:]:
        tx2 = tx2.replace(r.group(0), '', 1)
    return blocco[:i] + tx2 + blocco[j:]


def imposta_titolo(xml, testo):
    i = xml.find('<c:title>')
    if i < 0:
        return xml
    j = xml.find('</c:title>', i) + len('</c:title>')
    return xml[:i] + testo_singolo(xml[i:j], testo) + xml[j:]


def sostituisci(blocco, tag, dentro):
    """Rimpiazza <tag>...</tag> dentro una serie, inserendolo se manca."""
    m = re.search(r'<c:%s>.*?</c:%s>' % (tag, tag), blocco, flags=re.S)
    nuovo = '<c:%s>%s</c:%s>' % (tag, dentro, tag)
    if m:
        return blocco[:m.start()] + nuovo + blocco[m.end():]
    dove = blocco.find('<c:yVal>') if tag == 'xVal' else blocco.find('<c:smooth')
    return blocco[:dove] + nuovo + blocco[dove:]


class Dati:
    """Accesso ai valori del workbook, per ricostruire le cache dei grafici."""

    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path, data_only=True)

    def valori(self, foglio, col, r1, r2):
        ws = self.wb[foglio]
        c = col_index(col)
        return [ws.cell(r, c).value for r in range(r1, r2 + 1)]

    def num_ref(self, foglio, col, r1, r2):
        v = self.valori(foglio, col, r1, r2)
        pt = ''.join('<c:pt idx="%d"><c:v>%s</c:v></c:pt>' % (i, x)
                     for i, x in enumerate(v) if isinstance(x, (int, float)))
        cache = ('<c:numCache><c:formatCode>General</c:formatCode>'
                 '<c:ptCount val="%d"/>%s</c:numCache>' % (len(v), pt))
        return ("<c:numRef><c:f>'%s'!$%s$%d:$%s$%d</c:f>%s</c:numRef>"
                % (foglio, col, r1, col, r2, cache))

    def str_ref(self, foglio, col, riga, nome):
        return ("<c:strRef><c:f>'%s'!$%s$%d</c:f><c:strCache>"
                "<c:ptCount val=\"1\"/><c:pt idx=\"0\"><c:v>%s</c:v></c:pt>"
                "</c:strCache></c:strRef>" % (foglio, col, riga, esc(nome)))

    def data_set_completi(self):
        """Fogli 'Data Set N' con le stime compilate, in ordine."""
        out = []
        for nome in self.wb.sheetnames:
            m = re.match(r'Data Set (\d+)$', nome)
            if not m:
                continue
            ws = self.wb[nome]
            if all(ws.cell(r, col_index('K')).value is not None
                   for r in (2, 34, 66)):
                out.append(nome)
        return out

    def nodi_e_server(self, foglio, base):
        pop = self.wb[foglio].cell(base, col_index(COL_POP)).value
        nodi = int(round(pop / 100.0)) * 100
        return nodi, int(round(pop - nodi))


# --- ricostruzione di un grafico --------------------------------------------
def rigenera(xml, dati, foglio, base, etichetta, ds_label):
    """Riscrive titolo, serie e assi di un grafico sul blocco richiesto."""
    nodi, server = dati.nodi_e_server(foglio, base)
    titolo = '%s (%d nodi + %d server) - %s' % (etichetta, nodi, server, ds_label)
    xml = imposta_titolo(xml, titolo)
    serie = blocchi(xml, 'ser')       # dopo il titolo: le posizioni cambiano

    if len(serie) == len(SERIE_ACCUMULO) + 1:          # curve di accumulazione
        piano = [(nome, col, base + i * RIGHE_BLOCCO,
                  base + i * RIGHE_BLOCCO + RIGHE_BLOCCO - 1)
                 for i, (nome, col) in enumerate(SERIE_ACCUMULO)]
        piano.append(('popolazione_reale', COL_POP, base, base + RIGHE_BLOCCO - 1))
        nomi_da_cella = False
    elif len(serie) == len(SERIE_TRAFFICO) + 1:        # grafici per profilo
        r1, r2 = base, base + RIGHE_BLOCCO - 1
        piano = [(nome, col, r1, r2) for nome, col in SERIE_TRAFFICO]
        piano.append(('popolazione_reale', COL_POP, r1, r2))
        nomi_da_cella = True
    else:
        raise ValueError('grafico con %d serie: layout non riconosciuto'
                         % len(serie))

    # riga dell'intestazione del gruppo di nodi (1 per il gruppo che parte da 2)
    intestazione = [g for g in GRUPPI if g <= base <= g + 3 * RIGHE_BLOCCO][0] - 1

    pezzi, pos = [], 0
    for k, (nome, col, r1, r2) in enumerate(piano):
        b = serie[k].group(0)
        b = re.sub(r'<c:idx val="\d+"/>', '<c:idx val="%d"/>' % k, b)
        b = re.sub(r'<c:order val="\d+"/>', '<c:order val="%d"/>' % k, b)
        if nomi_da_cella:
            b = sostituisci(b, 'tx', dati.str_ref(foglio, col, intestazione, nome))
        else:
            b = sostituisci(b, 'tx', '<c:v>%s</c:v>' % esc(nome))
        b = sostituisci(b, 'xVal', dati.num_ref(foglio, COL_GIORNO, r1, r2))
        b = sostituisci(b, 'yVal', dati.num_ref(foglio, col, r1, r2))
        pezzi.append(xml[pos:serie[k].start()])
        pezzi.append(b)
        pos = serie[k].end()
    pezzi.append(xml[pos:])
    return sistema_assi(''.join(pezzi))


def raccogli_modelli_asse(cartella):
    """Memorizza un titolo asse ben formato per posizione (basso, sinistra)."""
    for f in sorted(os.listdir(cartella)):
        if not re.match(r'chart\d+\.xml$', f):
            continue
        xml = open(os.path.join(cartella, f), encoding='utf-8').read()
        for m in re.finditer(r'<c:valAx>.*?</c:valAx>', xml, flags=re.S):
            ax = m.group(0)
            p = re.search(r'<c:axPos val="(\w)"', ax).group(1)
            t = re.search(r'<c:title>.*?</c:title>', ax, flags=re.S)
            if t and re.search(r'<a:t>[^<]+</a:t>', t.group(0)) and p not in MODELLO_ASSE:
                MODELLO_ASSE[p] = t.group(0)
        if len(MODELLO_ASSE) == len(ASSI):
            return


def sistema_assi(xml):
    fuori, pos = [], 0
    for m in re.finditer(r'<c:valAx>.*?</c:valAx>', xml, flags=re.S):
        ax = m.group(0)
        p = re.search(r'<c:axPos val="(\w)"', ax).group(1)
        if p in ASSI:
            if '<c:majorGridlines' not in ax:
                g = re.search(r'<c:axPos val="\w"/>', ax)
                ax = ax[:g.end()] + '<c:majorGridlines/>' + ax[g.end():]
            t = re.search(r'<c:title>.*?</c:title>', ax, flags=re.S)
            usabile = t and re.search(r'<a:r>', t.group(0))
            if usabile:
                nuovo = testo_singolo(t.group(0), ASSI[p])
            elif p in MODELLO_ASSE:
                nuovo = testo_singolo(MODELLO_ASSE[p], ASSI[p])
            else:
                nuovo = None
            if nuovo is not None:
                if t:
                    ax = ax[:t.start()] + nuovo + ax[t.end():]
                else:
                    g = re.search(r'<c:minorGridlines/>|<c:majorGridlines/>', ax)
                    ax = ax[:g.end()] + nuovo + ax[g.end():]
        fuori.append(xml[pos:m.start()])
        fuori.append(ax)
        pos = m.end()
    fuori.append(xml[pos:])
    return ''.join(fuori)


# --- programma principale ----------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx', help='workbook da aggiornare')
    ap.add_argument('--out', help='file di destinazione (default: sovrascrive)')
    ap.add_argument('--no-pulizia', action='store_true',
                    help='non tocca la colonna Z ne\' la legenda colori')
    args = ap.parse_args()

    if not os.path.isfile(args.xlsx):
        sys.exit('File non trovato: %s' % args.xlsx)
    out = args.out or args.xlsx

    lavoro = tempfile.mkdtemp(prefix='grafici-')
    try:
        with zipfile.ZipFile(args.xlsx) as z:
            z.extractall(lavoro)
        costruisci(lavoro, args.xlsx, args.no_pulizia)

        tmp = out + '.tmp'
        with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as z:
            for radice, _, files in os.walk(lavoro):
                for f in files:
                    p = os.path.join(radice, f)
                    z.write(p, os.path.relpath(p, lavoro))
        os.replace(tmp, out)
    finally:
        shutil.rmtree(lavoro, ignore_errors=True)


def costruisci(lavoro, sorgente, salta_pulizia):
    grafici = os.path.join(lavoro, 'xl/charts')
    disegni = os.path.join(lavoro, 'xl/drawings')

    if not salta_pulizia:
        pulisci(lavoro)

    # i valori vanno letti dal file com'e' adesso, pulizia inclusa
    tmp = os.path.join(lavoro, '_dati.xlsx')
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as z:
        for radice, _, files in os.walk(lavoro):
            for f in files:
                if f == '_dati.xlsx':
                    continue
                p = os.path.join(radice, f)
                z.write(p, os.path.relpath(p, lavoro))
    dati = Dati(tmp)
    os.remove(tmp)

    raccogli_modelli_asse(grafici)
    fogli = dati.data_set_completi()
    if not fogli:
        sys.exit('Nessun foglio Data Set con le stime compilate.')

    etichette = etichette_per_disegno(lavoro)
    prossimo = 1 + max(int(re.match(r'chart(\d+)\.xml$', f).group(1))
                       for f in os.listdir(grafici)
                       if re.match(r'chart\d+\.xml$', f))
    usati = set()

    for dis, etichetta in sorted(etichette.items()):
        pd = os.path.join(disegni, 'drawing%d.xml' % dis)
        pr = os.path.join(disegni, '_rels/drawing%d.xml.rels' % dis)
        dxml = open(pd, encoding='utf-8').read()
        rxml = open(pr, encoding='utf-8').read()
        rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="\.\./charts/(chart\d+)\.xml"',
                               rxml))

        ancore = re.findall(r'<xdr:twoCellAnchor>.*?</xdr:twoCellAnchor>',
                            dxml, flags=re.S)
        modello = ancore[:3]           # la riga del primo Data Set
        sorgenti = [rels[re.search(r'r:id="(rId\d+)"', a).group(1)] for a in modello]

        nuove_ancore, nuove_rels, rid = [], [], 1
        oggetto = 1000 * dis
        for i, foglio in enumerate(fogli):
            shift = PASSO_RIGHE * i
            for k, sorg in enumerate(sorgenti):
                base_xml = open(os.path.join(grafici, '%s.xml' % sorg),
                                encoding='utf-8').read()
                base_riga = GRUPPI[k] if 'Catture' in etichetta else blocco_di(base_xml)
                if 'Catture' not in etichetta:
                    base_riga = blocco_di(base_xml)
                if i == 0:
                    n = int(re.match(r'chart(\d+)', sorg).group(1))
                else:
                    n = prossimo
                    prossimo += 1
                    vecchio = re.match(r'chart(\d+)', sorg).group(1)
                    for genere in ('colors', 'style'):
                        src_g = os.path.join(grafici, '%s%s.xml' % (genere, vecchio))
                        if os.path.isfile(src_g):
                            shutil.copy(src_g, os.path.join(grafici, '%s%d.xml' % (genere, n)))
                    rel_src = os.path.join(grafici, '_rels/%s.xml.rels' % sorg)
                    if os.path.isfile(rel_src):
                        s = open(rel_src, encoding='utf-8').read()
                        vecchio = re.match(r'chart(\d+)', sorg).group(1)
                        s = (s.replace('colors%s.xml' % vecchio, 'colors%d.xml' % n)
                              .replace('style%s.xml' % vecchio, 'style%d.xml' % n))
                        open(os.path.join(grafici, '_rels/chart%d.xml.rels' % n), 'w',
                             encoding='utf-8').write(s)

                nuovo = rigenera(base_xml, dati, foglio, base_riga, etichetta, foglio)
                open(os.path.join(grafici, 'chart%d.xml' % n), 'w',
                     encoding='utf-8').write(nuovo)
                usati.add(n)

                a = modello[k]
                a = re.sub(r'<xdr:row>(\d+)</xdr:row>',
                           lambda m: '<xdr:row>%d</xdr:row>' % (int(m.group(1)) + shift), a)
                a = re.sub(r'<a:extLst>.*?</a:extLst>', '', a, flags=re.S)
                oggetto += 1
                a = re.sub(r'<xdr:cNvPr id="\d+" name="[^"]*"',
                           '<xdr:cNvPr id="%d" name="Grafico %d"' % (oggetto, oggetto), a)
                a = re.sub(r'r:id="rId\d+"', 'r:id="rId%d"' % rid, a)
                nuove_ancore.append(a)
                nuove_rels.append(
                    '<Relationship Id="rId%d" Type="http://schemas.openxmlformats.org/'
                    'officeDocument/2006/relationships/chart" Target="../charts/chart%d.xml"/>'
                    % (rid, n))
                rid += 1

        testa = dxml[:dxml.find('<xdr:twoCellAnchor>')]
        open(pd, 'w', encoding='utf-8').write(testa + ''.join(nuove_ancore) + '</xdr:wsDr>')
        intestazione = rxml[:re.search(r'<Relationship\s', rxml).start()]
        open(pr, 'w', encoding='utf-8').write(
            intestazione + ''.join(nuove_rels) + '</Relationships>')

    scarta_avanzi(lavoro, usati)


def blocco_di(xml):
    """Prima riga del blocco dati a cui il grafico si riferisce."""
    return int(re.search(r'!\$[A-Z]+\$(\d+):', xml).group(1))


def etichette_per_disegno(lavoro):
    """{numero disegno: etichetta da usare nei titoli}, dal nome del foglio."""
    with open(os.path.join(lavoro, 'xl/workbook.xml'), encoding='utf-8') as f:
        wb = f.read()
    rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"',
                           open(os.path.join(lavoro, 'xl/_rels/workbook.xml.rels'),
                                encoding='utf-8').read()))
    out = {}
    for nome, rid in re.findall(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wb):
        percorso = rels[rid].split('/')[-1]
        p = os.path.join(lavoro, 'xl/worksheets', percorso)
        if not os.path.isfile(p):
            continue
        m = re.search(r'<drawing r:id="(rId\d+)"', open(p, encoding='utf-8').read())
        if not m:
            continue
        drel = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"',
                               open(os.path.join(lavoro, 'xl/worksheets/_rels',
                                                 percorso + '.rels'),
                                    encoding='utf-8').read()))
        num = int(re.search(r'drawing(\d+)\.xml', drel[m.group(1)]).group(1))
        etichetta = ('Traffico ' + nome.split(' - ', 1)[1]
                     if nome.startswith('Traffico - ') else 'Catture Dirette')
        out[num] = etichetta
    return out


def scarta_avanzi(lavoro, usati):
    """Elimina i grafici della corsa precedente e le voci in [Content_Types]."""
    grafici = os.path.join(lavoro, 'xl/charts')
    ct = os.path.join(lavoro, '[Content_Types].xml')
    testo = open(ct, encoding='utf-8').read()

    for f in os.listdir(grafici):
        m = re.match(r'(chart|colors|style)(\d+)\.xml$', f)
        if m and int(m.group(2)) not in usati:
            os.remove(os.path.join(grafici, f))
            testo = testo.replace(
                '<Override PartName="/xl/charts/%s" ContentType="%s"/>'
                % (f, tipo_di(m.group(1))), '')
    rels = os.path.join(grafici, '_rels')
    for f in (os.listdir(rels) if os.path.isdir(rels) else []):
        m = re.match(r'chart(\d+)\.xml\.rels$', f)
        if m and int(m.group(1)) not in usati:
            os.remove(os.path.join(rels, f))

    for n in sorted(usati):
        for genere in ('chart', 'colors', 'style'):
            voce = ('<Override PartName="/xl/charts/%s%d.xml" ContentType="%s"/>'
                    % (genere, n, tipo_di(genere)))
            if voce not in testo:
                testo = testo.replace('</Types>', voce + '</Types>')
    open(ct, 'w', encoding='utf-8').write(testo)


def tipo_di(genere):
    return {
        'chart': 'application/vnd.openxmlformats-officedocument.drawingml.chart+xml',
        'colors': 'application/vnd.ms-office.chartcolorstyle+xml',
        'style': 'application/vnd.ms-office.chartstyle+xml',
    }[genere]


def pulisci(lavoro):
    """Toglie la colonna Z (copia morta di X) e corregge la legenda colori."""
    fogli = os.path.join(lavoro, 'xl/worksheets')
    for f in os.listdir(fogli):
        if not re.match(r'sheet\d+\.xml$', f):
            continue
        p = os.path.join(fogli, f)
        s = open(p, encoding='utf-8').read()
        s = re.sub(r'<c r="Z\d+"(?:\s[^>]*)?(?:/>|>.*?</c>)', '', s, flags=re.S)
        s = re.sub(r'(<dimension ref="[A-Z]+\d+:)Z(\d+"/>)', r'\1X\2', s)
        open(p, 'w', encoding='utf-8').write(s)

    ss = os.path.join(lavoro, 'xl/sharedStrings.xml')
    if os.path.isfile(ss):
        s = open(ss, encoding='utf-8').read()
        s = s.replace('<t>err_rel &gt; 5%</t>', '<t>err_rel &gt; 20%</t>')
        open(ss, 'w', encoding='utf-8').write(s)


if __name__ == '__main__':
    main()