#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
grafici3d.py - grafici 3D dello spettro delle frequenze di cattura f_k
               a partire dal workbook di benchmark.

Assi:
    x = giorno di cattura (1..10)
    y = classe di frequenza k (1..5 e ">5")
    z = f_k, numero di nodi visti esattamente k volte

Non tocca il workbook: lo apre in sola lettura e scrive solo i PNG.

Di default elabora i fogli "Data Set 1..4"; "Data Set 5" viene saltato (si
disegna solo se richiesto con --sheet, o cambiando --escludi).

Esempi:
    python3 grafici3d.py --xlsx Benchmark_grafici.xlsx --tipo barre
    python3 grafici3d.py --tipo superficie --norm --out fig/
    python3 grafici3d.py --tipo waterfall --sheet "Data Set 1" --pop 1003
    python3 grafici3d.py --tipo barre --sheet "Data Set 1" --profilo bilanciato
    python3 grafici3d.py --tipo barre --escludi "Data Set 4" --escludi "Data Set 5"
"""

import argparse
import os

import numpy as np
import matplotlib
matplotlib.use("Agg")            # nessuna finestra: salva su file
import matplotlib.pyplot as plt
from matplotlib.collections import PolyCollection
from mpl_toolkits.mplot3d import Axes3D  # noqa: F401  (registra proj='3d')
import openpyxl

# --- layout del foglio (stesse colonne di bench_to_excel.py) -----------------
COL_PROFILO, COL_GIORNO, COL_T, COL_SOBS = 1, 2, 3, 4
COL_F1 = 5                       # E..I = f1..f5, J = f_piu_alte
COL_FN = 10
COL_POP = 24                     # X = popolazione_reale

CLASSI = ["1", "2", "3", "4", "5", ">5"]
PROFILI = ["molto_frequente", "bilanciato", "poco_frequente"]

# Fogli saltati quando non si passa --sheet. L'esclusione vale solo per la
# scoperta automatica: "--sheet 'Data Set 5'" lo disegna comunque.
FOGLI_ESCLUSI = {"data set 5"}


def leggi(xlsx, fogli=None, esclusi=FOGLI_ESCLUSI):
    """Ritorna una lista di record:
    {sheet, profilo, pop, giorni:[...], t:[...], sobs:[...], F: array (giorni x 6)}
    """
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    if fogli:
        nomi = fogli
    else:
        nomi = [s for s in wb.sheetnames
                if s.lower().startswith("data set") and s.lower() not in esclusi]
    blocchi = {}
    ordine = []

    for nome in nomi:
        ws = wb[nome]
        for riga in ws.iter_rows(min_row=1, max_col=COL_POP, values_only=True):
            prof = riga[COL_PROFILO - 1]
            giorno = riga[COL_GIORNO - 1]
            pop = riga[COL_POP - 1]
            if prof not in PROFILI or not isinstance(giorno, (int, float)):
                continue                       # intestazioni, righe vuote
            f = [riga[c - 1] for c in range(COL_F1, COL_FN + 1)]
            if any(v is None for v in f):
                continue                       # blocco non ancora popolato
            chiave = (nome, pop, prof)
            if chiave not in blocchi:
                blocchi[chiave] = {"sheet": nome, "pop": pop, "profilo": prof,
                                   "giorni": [], "t": [], "sobs": [], "F": []}
                ordine.append(chiave)
            b = blocchi[chiave]
            b["giorni"].append(int(giorno))
            b["t"].append(riga[COL_T - 1])
            b["sobs"].append(riga[COL_SOBS - 1])
            b["F"].append([float(v) for v in f])

    out = []
    for chiave in ordine:
        b = blocchi[chiave]
        idx = np.argsort(b["giorni"])          # giorni in ordine, per sicurezza
        b["giorni"] = np.array(b["giorni"])[idx]
        b["t"] = np.array(b["t"], dtype=float)[idx]
        b["sobs"] = np.array(b["sobs"], dtype=float)[idx]
        b["F"] = np.array(b["F"])[idx]
        out.append(b)
    return out


def normalizza(b):
    """f_k / S_obs: rende confrontabili popolazioni di taglia diversa."""
    s = np.where(b["sobs"] > 0, b["sobs"], 1.0)
    return b["F"] / s[:, None]


# --- i tre tipi di grafico ---------------------------------------------------
def barre(ax, x, y, Z, cmap="viridis"):
    """Istogramma 3D: una barra per ogni coppia (giorno, classe)."""
    X, Y = np.meshgrid(x, y, indexing="ij")
    xs, ys, zs = X.ravel(), Y.ravel(), Z.ravel()
    # colore per classe di frequenza: le file restano distinguibili di lato
    colori = plt.get_cmap(cmap)(np.linspace(0.15, 0.9, len(y)))
    c = np.tile(colori, (len(x), 1))
    ax.bar3d(xs - 0.35, ys - 0.35, np.zeros_like(zs),
             0.7, 0.7, zs, color=c, shade=True, edgecolor="k", linewidth=0.2)


def superficie(ax, x, y, Z, cmap="viridis"):
    """Superficie interpolata: mostra bene la forma dello spettro nel tempo."""
    X, Y = np.meshgrid(x, y, indexing="ij")
    ax.plot_surface(X, Y, Z, cmap=cmap, edgecolor="k", linewidth=0.3,
                    antialiased=True, rstride=1, cstride=1, alpha=0.95)
    ax.contour(X, Y, Z, zdir="z", offset=0, cmap=cmap, linewidths=0.6)


def waterfall(ax, x, y, Z, cmap="viridis"):
    """Un profilo pieno per giorno: buono quando i giorni vanno letti uno a uno."""
    colori = plt.get_cmap(cmap)(np.linspace(0.15, 0.9, len(x)))
    poligoni = []
    for i in range(len(x)):
        z = np.concatenate(([0.0], Z[i, :], [0.0]))
        yy = np.concatenate(([y[0]], y, [y[-1]]))
        poligoni.append(list(zip(yy, z)))
    pc = PolyCollection(poligoni, facecolors=colori, edgecolors="k",
                        linewidths=0.5, alpha=0.85)
    ax.add_collection3d(pc, zs=x, zdir="x")
    ax.set_xlim(x[0] - 0.5, x[-1] + 0.5)
    ax.set_ylim(y[0] - 0.5, y[-1] + 0.5)
    ax.set_zlim(0, float(Z.max()) * 1.05)


DISEGNA = {"barre": barre, "superficie": superficie, "waterfall": waterfall}


def figura(blocchi, tipo, norm, cmap, out, elev, azim):
    """Una figura per (foglio, popolazione): tre pannelli, uno per profilo."""
    b0 = blocchi[0]
    fig = plt.figure(figsize=(6.2 * len(blocchi), 5.0))
    zlab = "f_k / S_obs" if norm else "f_k  (nodi)"

    Zs = [normalizza(b) if norm else b["F"] for b in blocchi]
    zmax = max(float(Z.max()) for Z in Zs) * 1.05   # stessa scala nei pannelli

    for i, (b, Z) in enumerate(zip(blocchi, Zs), start=1):
        ax = fig.add_subplot(1, len(blocchi), i, projection="3d")
        x = b["giorni"].astype(float)
        y = np.arange(1, Z.shape[1] + 1, dtype=float)
        DISEGNA[tipo](ax, x, y, Z, cmap)

        ax.set_xlabel("giorno di cattura")
        ax.set_ylabel("classe di frequenza k")
        ax.set_zlabel(zlab)
        ax.set_yticks(y)
        ax.set_yticklabels(CLASSI)
        ax.set_xticks(x[::2])
        ax.set_zlim(0, zmax)
        ax.view_init(elev=elev, azim=azim)
        ax.set_title(b["profilo"].replace("_", " "), y=0.97)

    fig.suptitle("%s — N reale = %s — spettro delle frequenze di cattura"
                 % (b0["sheet"], b0["pop"]), fontsize=13, y=0.98)
    # tight_layout non sa gestire le decorazioni degli assi 3D: si regola a mano
    fig.subplots_adjust(left=0.02, right=0.98, top=0.88, bottom=0.02, wspace=0.05)

    os.makedirs(out, exist_ok=True)
    nome = "%s_pop%s_%s%s.png" % (b0["sheet"].replace(" ", ""), b0["pop"],
                                  tipo, "_norm" if norm else "")
    percorso = os.path.join(out, nome)
    fig.savefig(percorso, dpi=150)
    plt.close(fig)
    return percorso


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--xlsx", default="Benchmark_grafici.xlsx")
    p.add_argument("--sheet", action="append",
                   help="foglio da usare (ripetibile); default: i 'Data Set N' "
                        "tranne quelli in --escludi")
    p.add_argument("--escludi", action="append", metavar="FOGLIO",
                   help="foglio da saltare nella scoperta automatica "
                        "(ripetibile; default: 'Data Set 5')")
    p.add_argument("--pop", type=float, action="append",
                   help="filtra per popolazione reale (ripetibile)")
    p.add_argument("--profilo", action="append", choices=PROFILI,
                   help="filtra per profilo (ripetibile)")
    p.add_argument("--tipo", default="barre", choices=sorted(DISEGNA))
    p.add_argument("--norm", action="store_true",
                   help="normalizza f_k su S_obs")
    p.add_argument("--cmap", default="viridis")
    p.add_argument("--out", default="fig3d")
    p.add_argument("--elev", type=float, default=24.0)
    p.add_argument("--azim", type=float, default=-58.0)
    a = p.parse_args()

    esclusi = ({s.lower() for s in a.escludi} if a.escludi else FOGLI_ESCLUSI)
    dati = leggi(a.xlsx, a.sheet, esclusi)
    if a.pop:
        dati = [b for b in dati if b["pop"] in a.pop]
    if a.profilo:
        dati = [b for b in dati if b["profilo"] in a.profilo]
    if not dati:
        raise SystemExit("Nessun blocco corrisponde ai filtri.")

    # raggruppa per (foglio, popolazione) mantenendo l'ordine dei profili
    gruppi = {}
    for b in dati:
        gruppi.setdefault((b["sheet"], b["pop"]), []).append(b)
    for g in gruppi.values():
        g.sort(key=lambda b: PROFILI.index(b["profilo"]))

    for g in gruppi.values():
        print(figura(g, a.tipo, a.norm, a.cmap, a.out, a.elev, a.azim))


if __name__ == "__main__":
    main()